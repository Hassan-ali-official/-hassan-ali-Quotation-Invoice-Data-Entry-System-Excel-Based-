import tkinter
from tkinter import ttk
import os
import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.styles import Font, Color


def enter_data():
    filepath = "C:/Users/Hassan Ali/Desktop/python-mini-projects-master/projects/convert pdf to text/data.xlsx"
    company_name = nameEntry.get()
    ntn = ntnEntry.get()
    gst = GSTEntry.get()
    type_ = type1.get()
    date = date1.get()
    sr = sr_entry.get()
    desc = description_entry.get()
    dim = dimensions_entry.get()
    brand = brand_entry.get()
    qty_ = qty_entry.get()
    unit_ = unit_entry.get()
    price = price_entry.get()

    a = [company_name]
    b = [type_]
    c = ["ntn", ntn, " ", "gst", gst]
    d = ["Date", date]
    w = [sr, desc, dim, brand, qty, unit, price]
    q = ["Sr.", "Description", "dimmension", "Brand", "QTY", "Unit", "Price/unit", "Total price"]

    e = int(price) * (int(qty_))

    if not os.path.exists(filepath):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(a)
        sheet.append(b)
        sheet.append(c)
        sheet.append(d)
        sheet.append(q)
        workbook.save(filepath)
    workbook = openpyxl.load_workbook(filepath)
    ws = workbook['Sheet']

    ft = Font(color="FF0000", size=25)

    ws['A1'].font = ft
    sheet = workbook.active
    sheet.append([sr, desc, dim, brand, qty_, unit_, price, e, ])
    workbook.save(filepath)
    if sheet:
        print("good")


window = tkinter.Tk()  # parent window (widget) that contain all other widget
window.title("MAKE QUOTATIONS & INVOICES")  # CHANGE TITLE OF WINDOW
frame = tkinter.Frame(window)  # 2nd hierarchy frame inside of parent window
frame.pack()
# entering company information
company = tkinter.LabelFrame(frame, text="Enter company information")
company.grid(row=0, column=0)

companyName = tkinter.Label(company, text="Enter company name below")
companyName.grid(row=0, column=0)
nameEntry = ttk.Combobox(company, values=["ESK Enterprises", "Hassan Enterprises"])
nameEntry.grid(row=1, column=0)

companyNTN = tkinter.Label(company, text="Enter NTN number")
companyNTN.grid(row=0, column=1)
ntnEntry = tkinter.Entry(company)
ntnEntry.grid(row=1, column=1)

companyGST = tkinter.Label(company, text="Enter GST number")
companyGST.grid(row=0, column=2)
GSTEntry = tkinter.Entry(company)
GSTEntry.grid(row=1, column=2)
# padding
for widget1 in company.winfo_children():
    widget1.grid_configure(padx=10, pady=5)

# quotation/invoice/date
select = tkinter.LabelFrame(frame, text="Type and Date")
select.grid(row=1, column=0)
text = tkinter.Label(select, text="select type below")
text.grid(row=0, column=0)
type1 = ttk.Combobox(select, values=["Quotation", "Invoice"])
type1.grid(row=1, column=0, padx=20, pady=5)
date = tkinter.Label(select, text="DATE")
date.grid(row=0, column=1)
date1 = tkinter.Entry(select)
date1.grid(row=1, column=1, padx=20, pady=5)

# entering data
dataEntry = tkinter.LabelFrame(frame, text="Enter data below")
dataEntry.grid(row=3, column=0)

sr = tkinter.Label(dataEntry, text="Sr.")  # Serial number
sr.grid(row=0, column=0)
sr_entry = ttk.Spinbox(dataEntry, from_=1, to=100000)  # Serial number entry
sr_entry.grid(row=1, column=0)

description = tkinter.Label(dataEntry, text="Description")  # description
description.grid(row=0, column=1)
description_entry = tkinter.Entry(dataEntry)  # description entry
description_entry.grid(row=1, column=1)

dimensions = tkinter.Label(dataEntry, text="Dimensions")  # dimensions
dimensions.grid(row=0, column=2)
dimensions_entry = tkinter.Entry(dataEntry)  # dimensions entry
dimensions_entry.grid(row=1, column=2)

brand = tkinter.Label(dataEntry, text="brand")  # brand
brand.grid(row=0, column=3)
brand_entry = tkinter.Entry(dataEntry)  # brand entry
brand_entry.grid(row=1, column=3)

qty = tkinter.Label(dataEntry, text="qty")  # qty
qty.grid(row=0, column=4)
qty_entry = ttk.Spinbox(dataEntry, from_=1, to=100000000)  # qty entry
qty_entry.grid(row=1, column=4)

unit = tkinter.Label(dataEntry, text="unit")  # unit
unit.grid(row=0, column=5)
unit_entry = ttk.Combobox(dataEntry,
                          values=["PCS", "BOX", "MTR", "KM", "SQ_MTR", "FT", "SQ_FT", "INCHES"])  # unit entry
unit_entry.grid(row=1, column=5)

price = tkinter.Label(dataEntry, text="price/unit")  # price
price.grid(row=0, column=6)
price_entry = ttk.Spinbox(dataEntry, from_=1, to=500000000)  # price entry
price_entry.grid(row=1, column=6)
# padding
for widget2 in dataEntry.winfo_children():
    widget2.grid_configure(padx=10, pady=5)
# terms and condition
term = tkinter.LabelFrame(frame, text="TERMS & CONDITIONS")
term.grid(row=4, column=0)

term1 = tkinter.Label(term, text="Validity Duration")
term1.grid(row=0, column=0)

button2 = tkinter.Button(frame, text="ENTER DATA IN EXCEL", command=enter_data)
button2.grid(row=5, column=0, sticky="news", padx=25, pady=25)

window.mainloop()
